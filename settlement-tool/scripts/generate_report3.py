"""
결과물1(정산 요약본) -> 결과물3(세금계산서 요청용, 업체별 블록) 생성 스크립트

사용법:
    python generate_report3.py 결과물1.xlsx 결과물3_생성.xlsx

각 업체마다 화면(스크린샷)으로 캡처해서 메일에 붙이기 좋은 형태의
표 블록을 세로로 이어 붙인다. 시트 맨 위의 "작성일자" / "대표자명"
입력칸에 값을 채우면, 모든 업체 블록의 안내문구가 수식으로 자동 반영된다.
"""
import sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import FormulaRule

THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FONT = Font(name="맑은 고딕", bold=True, size=10)
BODY_FONT = Font(name="맑은 고딕", size=10)
NOTICE_FONT = Font(name="맑은 고딕", bold=True, color="FF0000", size=10)
INPUT_FONT = Font(name="맑은 고딕", bold=True, size=11, color="1155CC")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
NUM_FMT = '#,##0;-#,##0;"-"'

# 공급업체~매입금액을 붙여넣을 때 헷갈리지 않도록, 열 그룹마다 다른 연한
# 파스텔 색을 준다 (번호·업체명·구분 / 면세금액 / 과세금액(공급가액·부가세·
# 합계) / 매입금액). 헤더·데이터·소계 행 전부 같은 색 띠를 유지한다.
FILL_VENDOR = PatternFill("solid", fgColor="F2F2F2")    # 번호/업체명/구분 - 연회색
FILL_TAXFREE = PatternFill("solid", fgColor="E2EFDA")   # 면세금액 - 연녹색
FILL_TAXED = PatternFill("solid", fgColor="DDEBF7")     # 과세금액(공급가액/부가세/합계) - 연파랑
FILL_TOTAL = PatternFill("solid", fgColor="FFF2CC")     # 매입금액 - 연노랑


def col_fill(col):
    if col in (1, 2, 3):
        return FILL_VENDOR
    if col == 4:
        return FILL_TAXFREE
    if col in (5, 6, 7):
        return FILL_TAXED
    return FILL_TOTAL

BLOCK_HEIGHT = 7  # 헤더2 + 안내문구1 + 데이터1 + 공백1 + 소계1 + 여백1
DATE_CELL = "B1"
CEO_CELL = "B2"
CHECK_R3_CELL = "E5"   # 결과물3 안 소계 매입금액들의 합 (수식)
CHECK_R1_CELL = "E6"   # 결과물1 총지급예정액계 총계 (하이포스 합계행 값)
CHECK_DIFF_CELL = "E7"
CHECK_RESULT_CELL = "E8"
DATA_START_ROW = 11


def load_report1(path):
    raw = pd.read_excel(path)
    # 결과물1 맨 마지막 행은 하이포스가 붙이는 합계행(사업자번호가 비어있고
    # 거래처명 칸에 업체 수 숫자만 들어있다). 이 행의 총지급예정액계가
    # "진짜 총계"이므로 따로 떼어둔 다음, 업체 목록에서는 제외한다.
    grand_total = raw.loc[raw["사업자번호"].isna(), "총지급예정액계"].sum()
    df = raw[raw["사업자번호"].notna()]
    return df[["거래처명", "면세지급액", "과세공급가액", "부가세", "과세지급액", "총지급예정액계"]], grand_total


def style_range(ws, cell_range, border=True):
    for row in ws[cell_range]:
        for cell in row:
            if border:
                cell.border = BORDER


def write_input_header(ws, date_str=None, ceo_name=None):
    ws["A1"] = "작성일자"
    ws["A1"].font = HEADER_FONT
    ws["A1"].alignment = CENTER
    ws[DATE_CELL] = date_str or "7월 31일"
    ws[DATE_CELL].font = INPUT_FONT
    ws[DATE_CELL].fill = INPUT_FILL
    ws[DATE_CELL].alignment = Alignment(horizontal="center")

    ws["A2"] = "직매장 대표자명"
    ws["A2"].font = HEADER_FONT
    ws["A2"].alignment = CENTER
    ws[CEO_CELL] = ceo_name or "김건영"
    ws[CEO_CELL].font = INPUT_FONT
    ws[CEO_CELL].fill = INPUT_FILL
    ws[CEO_CELL].alignment = Alignment(horizontal="center")

    for c in ("A1", "A2"):
        ws[c].border = BORDER
    for c in (DATE_CELL, CEO_CELL):
        ws[c].border = BORDER

    note = ws.cell(row=3, column=1,
                    value="※ 위 두 칸만 바꾸면 아래 모든 업체 블록의 안내문구가 자동으로 바뀝니다.")
    note.font = Font(name="맑은 고딕", italic=True, size=9, color="666666")


def write_check_section(ws, subtotal_rows, grand_total):
    """아래 업체 블록들의 소계 매입금액 합과, 결과물1(전산원본) 합계행의
    총지급예정액계가 서로 맞는지 한눈에 보이는 검산 섹션을 상단에 둔다."""
    labels = [
        ("결과물3 합계 (소계 매입금액 합)", CHECK_R3_CELL,
         "=" + "+".join(f"H{r}" for r in subtotal_rows) if subtotal_rows else "=0"),
        ("결과물1 총계 (전산원본 합계행)", CHECK_R1_CELL, grand_total),
        ("차이", CHECK_DIFF_CELL, f"={CHECK_R3_CELL}-{CHECK_R1_CELL}"),
        ("검산 결과", CHECK_RESULT_CELL, f'=IF({CHECK_DIFF_CELL}=0,"일치","불일치 - 확인 필요")'),
    ]
    for i, (label, cell_ref, value) in enumerate(labels):
        row = 5 + i
        label_cell = ws.cell(row=row, column=1, value=label)
        label_cell.font = HEADER_FONT
        label_cell.alignment = Alignment(horizontal="left")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)

        val_cell = ws[cell_ref]
        val_cell.value = value
        val_cell.font = Font(name="맑은 고딕", bold=True, size=11)
        val_cell.alignment = Alignment(horizontal="center")
        if cell_ref != CHECK_RESULT_CELL:
            val_cell.number_format = NUM_FMT
        ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=8)
        for c in range(1, 9):
            ws.cell(row=row, column=c).border = BORDER

    ws.conditional_formatting.add(
        CHECK_RESULT_CELL,
        FormulaRule(formula=[f'{CHECK_RESULT_CELL}="일치"'], fill=PatternFill("solid", fgColor="C6EFCE"),
                    font=Font(color="1B6B34", bold=True)),
    )
    ws.conditional_formatting.add(
        CHECK_RESULT_CELL,
        FormulaRule(formula=[f'{CHECK_RESULT_CELL}<>"일치"'], fill=PatternFill("solid", fgColor="FFC7CE"),
                    font=Font(color="9C0006", bold=True)),
    )

    note = ws.cell(row=9, column=1,
                    value="※ 아래 업체별 소계를 전부 더한 값이 전산원본 합계행과 다르면, 업체가 빠졌거나 중복됐다는 뜻입니다.")
    note.font = Font(name="맑은 고딕", italic=True, size=9, color="666666")


def build_block(ws, start_row, seq, vendor_row):
    r = start_row
    name = vendor_row["거래처명"]
    tax_free = vendor_row["면세지급액"]
    supply = vendor_row["과세공급가액"]
    vat = vendor_row["부가세"]
    taxed_total = vendor_row["과세지급액"]
    grand_total = vendor_row["총지급예정액계"]

    # 헤더 1행: 공급업체 / 구분 / 면세금액 / 과세금액(병합) / 매입금액
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    ws.cell(row=r, column=1, value="공급업체").alignment = CENTER
    ws.cell(row=r, column=3, value="구분")
    ws.merge_cells(start_row=r, start_column=4, end_row=r + 1, end_column=4)
    ws.cell(row=r, column=4, value="면세금액")
    ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=7)
    ws.cell(row=r, column=5, value="과세금액")
    ws.merge_cells(start_row=r, start_column=8, end_row=r + 1, end_column=8)
    ws.cell(row=r, column=8, value="매입금액")

    # 헤더 2행: 과세금액 하위 라벨(공급가액/부가세/합계)
    ws.cell(row=r + 1, column=5, value="공급가액")
    ws.cell(row=r + 1, column=6, value="부가세")
    ws.cell(row=r + 1, column=7, value="합계")
    for c in range(1, 9):
        cell = ws.cell(row=r, column=c)
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.fill = col_fill(c)
        cell2 = ws.cell(row=r + 1, column=c)
        cell2.font = HEADER_FONT
        cell2.alignment = CENTER
        cell2.fill = col_fill(c)

    # 안내문구 (수식으로 작성일자/대표자명 반영)
    notice_row = r + 2
    ws.merge_cells(start_row=notice_row, start_column=1, end_row=notice_row, end_column=8)
    notice_cell = ws.cell(row=notice_row, column=1)
    notice_cell.value = (
        f'="공급받는자 : 춘천지역먹거리직매장 (대표자 "&${CEO_CELL}&")   /   "'
        f'&${DATE_CELL}&"자로 발행 해주세요."'
    )
    notice_cell.font = NOTICE_FONT
    notice_cell.alignment = Alignment(horizontal="center")

    # 업체 데이터 행
    data_row = r + 3
    ws.cell(row=data_row, column=1, value=seq)
    ws.cell(row=data_row, column=2, value=name)
    ws.cell(row=data_row, column=4, value=tax_free)
    ws.cell(row=data_row, column=5, value=supply)
    ws.cell(row=data_row, column=6, value=vat)
    ws.cell(row=data_row, column=7, value=taxed_total)
    ws.cell(row=data_row, column=8, value=grand_total)
    for c in range(1, 9):
        cell = ws.cell(row=data_row, column=c)
        cell.font = BODY_FONT
        cell.fill = col_fill(c)
        if c >= 4:
            cell.number_format = NUM_FMT
            cell.alignment = Alignment(horizontal="right")
        else:
            cell.alignment = CENTER if c != 2 else Alignment(horizontal="left", indent=1)

    # 공백 행
    blank_row = r + 4

    # 소계 행
    subtotal_row = r + 5
    ws.merge_cells(start_row=subtotal_row, start_column=1, end_row=subtotal_row, end_column=3)
    ws.cell(row=subtotal_row, column=1, value="소계")
    ws.cell(row=subtotal_row, column=4, value=f"=D{data_row}")
    ws.cell(row=subtotal_row, column=5, value=f"=E{data_row}")
    ws.cell(row=subtotal_row, column=6, value=f"=F{data_row}")
    ws.cell(row=subtotal_row, column=7, value=f"=G{data_row}")
    ws.cell(row=subtotal_row, column=8, value=f"=H{data_row}")
    for c in range(1, 9):
        cell = ws.cell(row=subtotal_row, column=c)
        cell.font = HEADER_FONT
        cell.fill = col_fill(c)
        if c >= 4:
            cell.number_format = NUM_FMT
            cell.alignment = Alignment(horizontal="right")
        else:
            cell.alignment = CENTER

    style_range(ws, f"A{r}:H{subtotal_row}")

    return r + BLOCK_HEIGHT, subtotal_row


def build_report3(df, grand_total, out_path, date_str=None, ceo_name=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "세금계산서 요청"

    write_input_header(ws, date_str, ceo_name)

    widths = [5, 22, 6, 11, 11, 9, 11, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    row = DATA_START_ROW
    seq = 1
    subtotal_rows = []
    for _, vendor_row in df.iterrows():
        if pd.isna(vendor_row["거래처명"]):
            continue
        row, subtotal_row = build_block(ws, row, seq, vendor_row)
        subtotal_rows.append(subtotal_row)
        seq += 1

    write_check_section(ws, subtotal_rows, grand_total)

    wb.save(out_path)
    print(f"결과물3 생성 완료: {out_path} ({seq - 1}개 업체 블록)")


if __name__ == "__main__":
    src, out = sys.argv[1], sys.argv[2]
    date_str = sys.argv[3] if len(sys.argv) > 3 else None
    ceo_name = sys.argv[4] if len(sys.argv) > 4 else None
    df, grand_total = load_report1(src)
    build_report3(df, grand_total, out, date_str, ceo_name)
